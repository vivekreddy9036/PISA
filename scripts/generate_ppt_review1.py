"""Generates the 1st-review presentation deck (docs/presentations/PISA_Review1.pptx).

Reuses the exact visual design system (colors, card/pill/header helpers)
from scripts/generate_ppt.py — the 0th-review deck — for visual continuity
between the two presentations. Content is entirely new: this deck covers
what was actually BUILT and VALIDATED since the 0th review (proposal
stage), not the literature-survey/proposal content of the earlier deck.

Every factual claim on these slides (test counts, real hardware results,
CVE data, architecture) is drawn directly from this project's actual
codebase and validation history — nothing here is invented for the slide.

Usage:
    venv/bin/python scripts/generate_ppt_review1.py
Output:
    docs/presentations/PISA_Review1.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─── COLOUR PALETTE (identical to scripts/generate_ppt.py) ─────────────────
BG_DARK       = RGBColor(0x09, 0x0D, 0x1A)
BG_CARD       = RGBColor(0x12, 0x18, 0x2E)
ACCENT_CYAN   = RGBColor(0x00, 0xD4, 0xFF)
ACCENT_GREEN  = RGBColor(0x00, 0xFF, 0x88)
ACCENT_RED    = RGBColor(0xFF, 0x4D, 0x4D)
ACCENT_ORANGE = RGBColor(0xFF, 0xA5, 0x00)
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_LIGHT    = RGBColor(0xB0, 0xB8, 0xC8)
GRAY_MED      = RGBColor(0x60, 0x6A, 0x80)
YELLOW        = RGBColor(0xFF, 0xE0, 0x66)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK_LAYOUT = prs.slide_layouts[6]

FOOTER_TEXT = "PISA — Portable IoT Security Assessment Platform  |  FYP Review 1  |  2026"


# ─── HELPERS (identical behavior to scripts/generate_ppt.py) ───────────────
def add_bg(slide, color=BG_DARK):
    bg = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = color; bg.line.fill.background()
    return bg


def add_rect(slide, l, t, w, h, color):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()
    return shape


def add_textbox(slide, text, l, t, w, h, font_size=18, bold=False, color=WHITE,
                 align=PP_ALIGN.LEFT, italic=False, wrap=True):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return box


def add_para(tf, text, font_size=16, bold=False, color=WHITE, align=PP_ALIGN.LEFT,
             space_before=0, italic=False):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return p


def slide_header(slide, title, subtitle=None, tag=None):
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.07), ACCENT_CYAN)
    add_textbox(slide, title, Inches(0.45), Inches(0.12), Inches(10), Inches(0.6),
                font_size=28, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.45), Inches(0.65), Inches(10), Inches(0.35),
                     font_size=14, color=ACCENT_CYAN)
    if tag:
        add_textbox(slide, tag, Inches(10.8), Inches(0.18), Inches(2.2), Inches(0.4),
                     font_size=11, color=GRAY_MED, align=PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(1, Inches(0.45), Inches(0.95), Inches(12.43), Inches(0.025))
    line.fill.solid(); line.fill.fore_color.rgb = GRAY_MED; line.line.fill.background()


def pill(slide, text, l, t, w, h, bg=ACCENT_CYAN, fg=BG_DARK, font_size=11, bold=True):
    add_rect(slide, l, t, w, h, bg)
    add_textbox(slide, text, l, t, w, h, font_size=font_size, bold=bold, color=fg, align=PP_ALIGN.CENTER)


def card(slide, l, t, w, h, accent_color=ACCENT_CYAN, title=None, body_lines=None,
         title_size=13, body_size=11):
    add_rect(slide, l, t, w, h, BG_CARD)
    add_rect(slide, l, t, Inches(0.06), h, accent_color)
    if title:
        add_textbox(slide, title, l + Inches(0.12), t + Inches(0.08), w - Inches(0.18), Inches(0.32),
                     font_size=title_size, bold=True, color=accent_color)
    if body_lines:
        tb = slide.shapes.add_textbox(l + Inches(0.12), t + Inches(0.38), w - Inches(0.2), h - Inches(0.48))
        tb.text_frame.word_wrap = True
        for i, line in enumerate(body_lines):
            p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
            p.space_before = Pt(2)
            run = p.add_run()
            run.text = line
            run.font.size = Pt(body_size)
            run.font.color.rgb = GRAY_LIGHT


def footer(slide, text=FOOTER_TEXT):
    add_rect(slide, 0, SLIDE_H - Inches(0.3), SLIDE_W, Inches(0.3), RGBColor(0x06, 0x0A, 0x14))
    add_textbox(slide, text, Inches(0.4), SLIDE_H - Inches(0.28), Inches(12.5), Inches(0.26),
                 font_size=9, color=GRAY_MED, align=PP_ALIGN.CENTER)


def new_slide():
    s = prs.slides.add_slide(BLANK_LAYOUT)
    add_bg(s)
    return s


# =============================================================================
# SLIDE 1 — TITLE
# =============================================================================
s = new_slide()
for i in range(20):
    line = s.shapes.add_shape(1, 0, Inches(i * 0.4), SLIDE_W, Inches(0.01))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(0x14, 0x1C, 0x30); line.line.fill.background()
add_rect(s, 0, 0, SLIDE_W, Inches(0.1), ACCENT_CYAN)
add_rect(s, 0, 0, Inches(0.1), SLIDE_H, ACCENT_CYAN)
add_rect(s, Inches(0.7), Inches(1.1), Inches(3.0), Inches(1.8), BG_CARD)
add_rect(s, Inches(0.7), Inches(1.1), Inches(0.12), Inches(1.8), ACCENT_GREEN)
add_textbox(s, "PISA", Inches(0.9), Inches(1.15), Inches(2.6), Inches(1.0),
            font_size=72, bold=True, color=ACCENT_GREEN)
add_textbox(s, "1st Review  |  Implementation & Validation", Inches(0.9), Inches(2.05), Inches(2.6), Inches(0.4),
            font_size=11, color=GRAY_MED)
add_textbox(s, "Portable IoT Security Assessment Platform", Inches(0.7), Inches(3.1), Inches(11.0), Inches(0.8),
            font_size=36, bold=True, color=WHITE)
add_textbox(s, "From Proposal to Working System — M0 through M5 Implemented, Tested (380 Passing Tests),\n"
               "and Validated Against Real Hardware and Live Vulnerability Intelligence APIs",
            Inches(0.7), Inches(3.95), Inches(11.0), Inches(0.9), font_size=16, color=GRAY_LIGHT)
add_rect(s, Inches(0.7), Inches(4.95), Inches(11.0), Inches(0.04), ACCENT_GREEN)
pill(s, "Real AWUS036ACM Hardware", Inches(0.7), Inches(5.1), Inches(2.6), Inches(0.38),
     bg=RGBColor(0x00, 0x3A, 0x26), fg=ACCENT_GREEN, font_size=11)
pill(s, "380 Tests Passing", Inches(3.45), Inches(5.1), Inches(2.0), Inches(0.38),
     bg=RGBColor(0x00, 0x3A, 0x52), fg=ACCENT_CYAN, font_size=11)
pill(s, "Live NVD / EPSS / KEV", Inches(5.6), Inches(5.1), Inches(2.4), Inches(0.38),
     bg=RGBColor(0x3A, 0x28, 0x00), fg=ACCENT_ORANGE, font_size=11)
pill(s, "Authorized Use Only", Inches(8.15), Inches(5.1), Inches(2.4), Inches(0.38),
     bg=RGBColor(0x3A, 0x00, 0x00), fg=ACCENT_RED, font_size=11)
add_textbox(s, "Vivek Reddy  |  20CYS495  |  Amrita School of Engineering",
            Inches(0.7), Inches(5.7), Inches(11.0), Inches(0.4), font_size=13, color=GRAY_LIGHT)

# =============================================================================
# SLIDE 2 — AGENDA
# =============================================================================
s = new_slide(); slide_header(s, "Agenda", tag="Review 1"); footer(s)
items = [
    ("01", "Recap", "What was approved at the 0th review"),
    ("02", "Architecture", "The full M0–M5 pipeline as actually built"),
    ("03", "What Makes PISA Different", "vs. Nmap + Bettercap + manual CVE lookup + RouterSploit"),
    ("04", "Module Walkthrough", "M0 WiFi → M1 Discovery → M2 Fingerprint → M3 Vuln Intel"),
    ("05", "Security Architecture", "Applicability → Verification → Authorization → Exploitation gate"),
    ("06", "Real Hardware Validation", "What has actually been proven, on real evidence"),
    ("07", "Engineering Rigor", "380 automated tests, honest known limitations"),
    ("08", "Roadmap", "What's next before the final review"),
]
W, H, GAP_X, GAP_Y, SX, SY = Inches(5.9), Inches(0.72), Inches(0.25), Inches(0.12), Inches(0.45), Inches(1.15)
for i, (num, title, desc) in enumerate(items):
    lx, ty = SX + (i % 2) * (W + GAP_X), SY + (i // 2) * (H + GAP_Y)
    add_rect(s, lx, ty, W, H, BG_CARD)
    add_rect(s, lx, ty, Inches(0.55), H, ACCENT_GREEN)
    add_textbox(s, num, lx, ty, Inches(0.55), H, font_size=20, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER)
    add_textbox(s, title, lx + Inches(0.62), ty + Inches(0.06), W - Inches(0.7), Inches(0.32), font_size=15, bold=True, color=WHITE)
    add_textbox(s, desc, lx + Inches(0.62), ty + Inches(0.36), W - Inches(0.7), Inches(0.3), font_size=10, color=GRAY_LIGHT)

# =============================================================================
# SLIDE 3 — RECAP OF 0TH REVIEW
# =============================================================================
s = new_slide(); slide_header(s, "Recap", subtitle="What Was Approved at the 0th Review", tag="01 / 08"); footer(s)
card(s, Inches(0.45), Inches(1.15), Inches(6.0), Inches(2.5), ACCENT_CYAN, "The Approved Problem",
     ["IoT security assessment today needs 4+ separate tools (Bettercap, Nmap, manual",
      "NVD lookup, RouterSploit), each with no shared identity or evidence model.",
      "", "No existing tool combines WiFi-layer posture scoring with a full device-level",
      "vulnerability pipeline in one portable package."], body_size=12)
card(s, Inches(6.65), Inches(1.15), Inches(6.2), Inches(2.5), ACCENT_GREEN, "The Approved Scope (v1)",
     ["M0: Passive WiFi beacon capture + WSPS security scoring",
      "M1: Network join + ARP/Nmap device discovery",
      "M2: Protocol behavioral fingerprinting (HTTP/MQTT/CoAP/RTSP)",
      "M3: CPE-precise CVE correlation + CVSS/EPSS/KEV scoring",
      "M4: Authorized, gated exploit verification"], body_size=12)
card(s, Inches(0.45), Inches(3.85), Inches(12.4), Inches(2.5), ACCENT_ORANGE, "This Review's Question",
     ["Everything approved in the proposal has now been implemented, tested (380 automated",
      "tests), and — for the first time — validated against real physical hardware: a real",
      "monitor-mode WiFi adapter, real beacon capture, real network discovery, and real, live",
      "calls to NVD/FIRST.org/CISA. This review demonstrates that system, honestly, including",
      "what it has and hasn't found real vulnerabilities on yet."], body_size=13)

# =============================================================================
# SLIDE 4 — SYSTEM ARCHITECTURE
# =============================================================================
s = new_slide(); slide_header(s, "System Architecture", subtitle="The Full M0–M5 Pipeline, As Actually Built", tag="02 / 08"); footer(s)
stages = [
    ("M0", "WiFi Assessment", "Beacon capture (Scapy, monitor mode) -> WSPS scoring (7-factor formula, A-F grade)"),
    ("M1", "Network Discovery", "WiFi join (nmcli) -> ARP sweep (arp-scan) -> Nmap port/OS scan -> mDNS ID"),
    ("M2", "Protocol Fingerprinting", "HTTP / MQTT / CoAP / RTSP probes -> device type + structured identity fusion"),
    ("M3", "Vulnerability Intelligence", "CPE mapping (live NVD) -> CVE correlation -> CVSS+EPSS+KEV -> Applicability"),
    ("M4", "Verification & Exploitation", "Safe live verification -> explicit authorization -> gated RouterSploit exploit"),
    ("M5", "Dashboard / API", "Flask web UI + REST API surfacing every stage above, loopback-only by default"),
]
y = Inches(1.2)
for code, title, desc in stages:
    add_rect(s, Inches(0.45), y, Inches(1.0), Inches(0.85), ACCENT_GREEN)
    add_textbox(s, code, Inches(0.45), y + Inches(0.18), Inches(1.0), Inches(0.5), font_size=24, bold=True,
                color=BG_DARK, align=PP_ALIGN.CENTER)
    card(s, Inches(1.6), y, Inches(11.25), Inches(0.85), ACCENT_CYAN, title, [desc], title_size=14, body_size=11)
    y += Inches(0.97)

# =============================================================================
# SLIDE 5 — WHAT MAKES PISA DIFFERENT
# =============================================================================
s = new_slide(); slide_header(s, "What Makes PISA Different", subtitle="vs. the Current State of the Art", tag="03 / 08"); footer(s)
rows = [
    ("Capability", "Nmap / Bettercap / manual", "RouterSploit alone", "PISA"),
    ("WiFi posture scoring", "No", "No", "Yes — 7-factor WSPS, A-F grade"),
    ("CPE-precise CVE match", "No (keyword search only)", "No", "Yes — real NVD CPE Dictionary, never fabricated"),
    ("Applicability reasoning", "No", "No", "Yes — AND/OR/version-range config evaluation"),
    ("Safe live verification", "No", "No", "Yes — before any exploit is even considered"),
    ("Gated authorization", "N/A", "Manual, no state machine", "Server-side: AFFECTED + VERIFIED required"),
    ("Single shared evidence model", "No — 4 disconnected tools", "No", "Yes — one identity, one DB, one pipeline"),
    ("Portable / edge-deployable", "Partial", "Partial", "Yes — Raspberry Pi + external radio, no cloud required"),
]
col_w = [Inches(3.0), Inches(3.3), Inches(2.7), Inches(3.65)]
x = Inches(0.45); y = Inches(1.15)
for j, (a, b, c, d) in enumerate([rows[0]]):
    xx = x
    for k, val in enumerate([a, b, c, d]):
        add_rect(s, xx, y, col_w[k], Inches(0.4), RGBColor(0x1a, 0x24, 0x3c))
        add_textbox(s, val, xx + Inches(0.08), y + Inches(0.05), col_w[k] - Inches(0.16), Inches(0.32),
                    font_size=11, bold=True, color=ACCENT_CYAN)
        xx += col_w[k]
y += Inches(0.4)
for row in rows[1:]:
    xx = x
    bg = BG_CARD
    for k, val in enumerate(row):
        add_rect(s, xx, y, col_w[k], Inches(0.55), bg)
        color = ACCENT_GREEN if k == 3 else GRAY_LIGHT
        weight = k == 3
        add_textbox(s, val, xx + Inches(0.08), y + Inches(0.08), col_w[k] - Inches(0.16), Inches(0.42),
                    font_size=10, bold=weight, color=color)
        xx += col_w[k]
    y += Inches(0.57)

# =============================================================================
# SLIDE 6 — M0 WIFI ASSESSMENT
# =============================================================================
s = new_slide(); slide_header(s, "M0 — WiFi Assessment", subtitle="Real Hardware, Real Beacon Capture, Real Scoring", tag="04 / 08"); footer(s)
card(s, Inches(0.45), Inches(1.15), Inches(6.0), Inches(2.9), ACCENT_CYAN, "WSPS Formula (7 factors, additive, 0-100)",
     ["Encryption: WPA3 +35 / WPA2 +25 / WPA +10 / Open +0",
      "Channel: non-overlapping (1/6/11) +15, other +5",
      "Signal strength: >-50dBm +20, >-70 +12, >-85 +6",
      "Beacon interval == 100ms: +12",
      "Hidden SSID: -10   |   PMF enabled: +15   |   WPS enabled: -15",
      "Grade thresholds: A>=90, B>=75, C>=60, D>=45, E>=30, else F"], body_size=12)
card(s, Inches(6.65), Inches(1.15), Inches(6.2), Inches(2.9), ACCENT_GREEN, "REAL Hardware Result (this session)",
     ["Adapter: Alfa AWUS036ACM (MediaTek MT7612U, mt76x2u driver)",
      "Real monitor mode confirmed via `iw dev ... info`",
      "Real 802.11 beacon frames received and parsed",
      "4 real networks captured on live campus WiFi:",
      "  Amrita / Amrita_CHN2 / Amrita_CH / Amrita_Guest",
      "All WPA2, ch36, real RSSI -55 to -57dBm -> WSPS Grade D (54)",
      "Result persisted to a real SQLite session row"], body_size=12)
card(s, Inches(0.45), Inches(4.25), Inches(12.4), Inches(1.9), ACCENT_ORANGE, "Classification",
     ["PASS -- REAL HARDWARE. Not simulated, not mocked -- genuine RF reception through a",
      "USB monitor-mode radio, parsed by the actual pisa/m0/beacon_capture.py implementation,",
      "with the score computed by the real, unmodified WSPS function."], body_size=13)

# =============================================================================
# SLIDE 7 — M1 NETWORK DISCOVERY
# =============================================================================
s = new_slide(); slide_header(s, "M1 — Network Discovery", subtitle="Real ARP + Nmap Against a Real Live Subnet", tag="04 / 08"); footer(s)
card(s, Inches(0.45), Inches(1.15), Inches(6.0), Inches(2.9), ACCENT_CYAN, "How It Works",
     ["nmcli joins the target WiFi network (managed-mode radio)",
      "arp-scan sweeps the joined subnet for live hosts",
      "Nmap (-sV -O) scans each host for IoT-relevant open ports",
      "  (22,23,80,443,502,554,1883,5555,5683,8080,8443)",
      "mDNS queries add self-announced device names where available",
      "Every result persists to SQLite via real query functions"], body_size=12)
card(s, Inches(6.65), Inches(1.15), Inches(6.2), Inches(2.9), ACCENT_GREEN, "Real Single-Target Validation",
     ["Target: a real, operator-owned device on the campus network",
      "Real ARP resolution: genuine MAC found, even though ICMP was blocked",
      "Real Nmap -sV -O: ran as root, ~13s real network round-trip",
      "Real OUI vendor lookup against the live IEEE registry",
      "Honest result: 0 open IoT-relevant ports -- correct, since the",
      "target was a general-purpose laptop, not an IoT device"], body_size=12)
card(s, Inches(0.45), Inches(4.25), Inches(12.4), Inches(1.9), ACCENT_ORANGE, "A Deliberate Scope Decision",
     ["A full subnet-wide scan was investigated and intentionally NOT run live during this",
      "project's validation -- the current implementation has no target-range restriction, and",
      "scanning an entire shared campus subnet raises real authorization-scope questions that",
      "were treated seriously rather than glossed over. Single, explicitly authorized targets",
      "were used instead."], body_size=12)

# =============================================================================
# SLIDE 8 — M2 PROTOCOL FINGERPRINTING
# =============================================================================
s = new_slide(); slide_header(s, "M2 — Protocol Fingerprinting", subtitle="Real Sockets, Real Banners, Real Structured Identity", tag="04 / 08"); footer(s)
card(s, Inches(0.45), Inches(1.15), Inches(6.0), Inches(2.9), ACCENT_CYAN, "Four Real Probes",
     ["HTTP: GET /, parses Server header + keyword match",
      "MQTT: anonymous CONNACK -- any reply proves a broker",
      "CoAP: GET /.well-known/core -- resource enumeration",
      "RTSP: raw OPTIONS request -- camera-signature status line",
      "Fusion: additive confidence sum across evidence, argmax wins",
      "No ML, no black box -- every score traces to a real, named rule"], body_size=12)
card(s, Inches(6.65), Inches(1.15), Inches(6.2), Inches(2.9), ACCENT_GREEN, "Real Controlled-Target Result",
     ["Target: a local, controlled demo IoT service (real sockets,",
      "real HTTP/CoAP/RTSP responses -- clearly not a commercial device)",
      "Real result: device_type = \"IP Camera\", confidence 1.0",
      "Real structured identity extracted from real banners:",
      "  product=\"IPCamera-WebServer\", version=\"1.0\"",
      "5 real evidence rows persisted to fingerprint_signatures"], body_size=12)
card(s, Inches(0.45), Inches(4.25), Inches(12.4), Inches(1.9), ACCENT_ORANGE, "Classification",
     ["PASS -- REAL NETWORK / REAL PISA M2. Genuine TCP/UDP traffic and genuine protocol",
      "parsing, run through the unmodified pisa/m2 probes -- the controlled target's synthetic",
      "banners are clearly labeled as such wherever shown, never presented as a real product."], body_size=13)

# =============================================================================
# SLIDE 9 — M3 VULNERABILITY INTELLIGENCE
# =============================================================================
s = new_slide(); slide_header(s, "M3 — Vulnerability Intelligence", subtitle="CPE -> CVE -> CVSS/EPSS/KEV -> Applicability, All Live", tag="04 / 08"); footer(s)
card(s, Inches(0.45), Inches(1.15), Inches(6.0), Inches(2.9), ACCENT_CYAN, "The Pipeline",
     ["Identity -> live NVD CPE Dictionary keyword search",
      "CPE match -> live NVD CVE API 2.0, paginated + retried",
      "Every finding enriched with live EPSS (FIRST.org) + CISA KEV",
      "ExploitScore = 40% CVSS + 40% EPSS + 20% KEV, one number",
      "Applicability engine evaluates real AND/OR/version-range",
      "  configuration trees -- never assumes, only concludes"], body_size=12)
card(s, Inches(6.65), Inches(1.15), Inches(6.2), Inches(2.9), ACCENT_GREEN, "Real Live Results, Both Directions",
     ["Controlled target: real NVD query for a synthetic banner",
      "  -> correctly, honestly NO_CPE_DATA (no fake match invented)",
      "Reference CVE (Xiongmai CVE-2017-7577): real CPE match,",
      "  real CVSS 9.8, real EPSS 29%, real applicability = AFFECTED",
      "Applicability states are never guessed:",
      "  AFFECTED / POTENTIALLY_AFFECTED / UNKNOWN / NOT_APPLICABLE"], body_size=12)
card(s, Inches(0.45), Inches(4.25), Inches(12.4), Inches(1.9), ACCENT_ORANGE, "The Core Engineering Principle",
     ["\"CVE exists\" is never presented as \"device is vulnerable.\" A keyword-only match is",
      "structurally incapable of reaching AFFECTED in this codebase -- enforced in code, proven",
      "by a dedicated automated test, not just claimed in a slide."], body_size=13)

# =============================================================================
# SLIDE 10 — SECURITY ARCHITECTURE / STATE MACHINE
# =============================================================================
s = new_slide(); slide_header(s, "Security Architecture", subtitle="The Applicability -> Verification -> Authorization -> Exploitation Gate", tag="05 / 08"); footer(s)
gate_stages = [
    ("CVE Found", "device_cves row exists -- correlation only, no claim of vulnerability", ACCENT_CYAN),
    ("Applicability", "AFFECTED required (or POTENTIALLY_AFFECTED for a supporting test)", ACCENT_CYAN),
    ("Verification", "A real, safe, live check must return VERIFIED_VULNERABLE", ACCENT_ORANGE),
    ("Authorization", "A named human operator must explicitly authorize -- never automatic", ACCENT_ORANGE),
    ("Exploitation", "Only then does the registered, gated RouterSploit adapter run", ACCENT_RED),
]
y = Inches(1.2)
for title, desc, color in gate_stages:
    card(s, Inches(0.45), y, Inches(12.4), Inches(0.85), color, title, [desc], title_size=15, body_size=12)
    y += Inches(0.95)

# =============================================================================
# SLIDE 11 — PROVEN NEGATIVE CONTROLS
# =============================================================================
s = new_slide(); slide_header(s, "The Gate Fails Closed — Proven, Not Claimed", subtitle="Live HTTP Checks Against the Real Running Dashboard", tag="05 / 08"); footer(s)
checks = [
    ("NOT_APPLICABLE finding", "403 Forbidden, gate_state=BLOCKED_APPLICABILITY, zero exploit attempts recorded"),
    ("AFFECTED but NOT_VERIFIED", "403 Forbidden, gate_state=BLOCKED_VERIFICATION"),
    ("Missing authorization", "400 Bad Request, even on a fully eligible finding"),
    ("Client-supplied arbitrary module path", "Silently ignored -- the real registry path executes instead, never the client's"),
    ("A genuine RouterSploit timeout", "Persisted as its own TIMEOUT status, never disguised as a generic failure"),
]
y = Inches(1.2)
for title, desc in checks:
    add_rect(s, Inches(0.45), y, Inches(0.5), Inches(0.75), ACCENT_RED)
    add_textbox(s, "X", Inches(0.45), y + Inches(0.18), Inches(0.5), Inches(0.4), font_size=18, bold=True,
                color=WHITE, align=PP_ALIGN.CENTER)
    card(s, Inches(1.1), y, Inches(11.75), Inches(0.75), ACCENT_RED, title, [desc], title_size=13, body_size=11)
    y += Inches(0.85)

# =============================================================================
# SLIDE 12 — REAL HARDWARE VALIDATION SUMMARY
# =============================================================================
s = new_slide(); slide_header(s, "Real Hardware Validation Summary", subtitle="What Has Actually Been Proven, on Real Evidence", tag="06 / 08"); footer(s)
table = [
    ("Module", "Status", "Evidence"),
    ("M0 WiFi Assessment", "PASS", "Real AWUS036ACM, real monitor mode, real beacon capture, real WSPS score"),
    ("M1 Network Discovery", "PASS", "Real ARP resolution + real Nmap -sV -O against a real, authorized target"),
    ("M2 Fingerprinting", "PASS", "Real protocol probes against a real (controlled) service, real evidence"),
    ("M3 Vulnerability Intel", "PASS", "Real, live NVD/EPSS/KEV calls -- both a real match and an honest non-match"),
    ("M4 Exploitation Gate", "PASS (software)", "Gate proven via live HTTP negative controls -- physical exploit target pending"),
    ("M5 Dashboard", "PASS", "Real Flask app, all routes verified, real + demo modes both rendering correctly"),
]
col_w = [Inches(3.3), Inches(2.2), Inches(7.1)]
x, y = Inches(0.45), Inches(1.2)
for i, row in enumerate(table):
    xx = x
    header = i == 0
    for k, val in enumerate(row):
        bgc = RGBColor(0x1a, 0x24, 0x3c) if header else BG_CARD
        add_rect(s, xx, y, col_w[k], Inches(0.55), bgc)
        color = ACCENT_CYAN if header else (ACCENT_GREEN if k == 1 else GRAY_LIGHT)
        add_textbox(s, val, xx + Inches(0.1), y + Inches(0.08), col_w[k] - Inches(0.18), Inches(0.4),
                     font_size=11, bold=header or k == 1, color=color)
        xx += col_w[k]
    y += Inches(0.57)
card(s, Inches(0.45), y + Inches(0.15), Inches(12.4), Inches(1.0), ACCENT_ORANGE, "Honest Gap",
     ["No real, physically vulnerable IoT device has been tested yet -- every real M3 result so far has been an",
      "honest non-match. Acquiring the already-identified target hardware is the top item on the roadmap."], body_size=12)

# =============================================================================
# SLIDE 13 — ENGINEERING RIGOR / TESTS
# =============================================================================
s = new_slide(); slide_header(s, "Engineering Rigor", subtitle="Automated Test Coverage & Honest Known Limitations", tag="07 / 08"); footer(s)
add_rect(s, Inches(0.45), Inches(1.2), Inches(3.6), Inches(1.5), BG_CARD)
add_textbox(s, "380", Inches(0.45), Inches(1.3), Inches(3.6), Inches(0.9), font_size=54, bold=True,
            color=ACCENT_GREEN, align=PP_ALIGN.CENTER)
add_textbox(s, "tests passing, 1 skipped, 0 failed", Inches(0.45), Inches(2.15), Inches(3.6), Inches(0.4),
            font_size=12, color=GRAY_LIGHT, align=PP_ALIGN.CENTER)
card(s, Inches(4.25), Inches(1.2), Inches(8.6), Inches(1.5), ACCENT_CYAN, "Coverage by Module",
     ["M0: 39  |  M1: 30  |  M2: 41  |  M3: 179  |  M4: 18  |  M5: 32  |  DB: 42",
      "M3's vulnerability-intelligence pipeline carries the most tests -- it's both the largest",
      "and the most safety-critical part of the system."], body_size=12)
card(s, Inches(0.45), Inches(2.95), Inches(12.4), Inches(2.1), ACCENT_ORANGE, "Known Limitations — Stated Proactively",
     ["No real vulnerable IoT device has been physically validated yet (M0/M1/M2 hardware is real; the specific",
      "device that would trigger a real positive CVE match is the next acquisition, already scoped).",
      "M1's discovery has no built-in target-range restriction -- a real, acknowledged architecture note.",
      "RouterSploit's execution timeout is thread-bound, not process-isolated -- a documented, known limitation.",
      "The demo mode uses clearly-labeled seeded data to visualize the full pipeline -- never presented as real."],
     body_size=12)

# =============================================================================
# SLIDE 14 — DEMO MODE
# =============================================================================
s = new_slide(); slide_header(s, "Live Demonstration Plan", subtitle="What Will Be Shown, and How", tag="Live Demo"); footer(s)
demo_steps = [
    ("1", "Real M0", "AWUS036ACM in monitor mode -- live beacon capture -- real WSPS grades on screen"),
    ("2", "Real M1/M2/M3 evidence", "Cite the real single-target validation -- real ARP/Nmap/NVD calls, honest result"),
    ("3", "Controlled local target", "Real HTTP/RTSP/CoAP service -- real M2 fingerprint -- real (honest) M3 lookup"),
    ("4", "DEMO mode", "Clearly labeled SIMULATED/REPLAY dashboard -- full pipeline visualization, all 4 real states"),
    ("5", "Security gate", "Live API calls proving NOT_APPLICABLE / NOT_VERIFIED / no-auth all fail closed"),
    ("6", "Test suite", "pytest tests/ -v -- 380 passed, on screen, live"),
]
y = Inches(1.2)
for num, title, desc in demo_steps:
    add_rect(s, Inches(0.45), y, Inches(0.6), Inches(0.72), ACCENT_GREEN)
    add_textbox(s, num, Inches(0.45), y + Inches(0.15), Inches(0.6), Inches(0.45), font_size=22, bold=True,
                color=BG_DARK, align=PP_ALIGN.CENTER)
    card(s, Inches(1.15), y, Inches(11.7), Inches(0.72), ACCENT_CYAN, title, [desc], title_size=13, body_size=11)
    y += Inches(0.8)

# =============================================================================
# SLIDE 15 — ROADMAP
# =============================================================================
s = new_slide(); slide_header(s, "Roadmap", subtitle="What's Next Before the Final Review", tag="08 / 08"); footer(s)
roadmap = [
    ("A", "Physical vulnerable-target acquisition", "Xiongmai-family device -- already scoped, live-verified CPE match potential"),
    ("B", "Full positive M1-M4 chain on real hardware", "CPE match -> AFFECTED -> VERIFIED_VULNERABLE -> authorized exploitation"),
    ("C", "Raspberry Pi edge deployment", "Move from dev laptop to the target field form factor"),
    ("D", "Registry expansion", "Add further vetted, tested CVE modules beyond the current two"),
    ("E", "Final report & demo rehearsal", "Full write-up + live run-through ahead of the final review"),
]
y = Inches(1.2)
for letter, title, desc in roadmap:
    add_rect(s, Inches(0.45), y, Inches(0.6), Inches(0.85), ACCENT_ORANGE)
    add_textbox(s, letter, Inches(0.45), y + Inches(0.2), Inches(0.6), Inches(0.5), font_size=22, bold=True,
                color=BG_DARK, align=PP_ALIGN.CENTER)
    card(s, Inches(1.15), y, Inches(11.7), Inches(0.85), ACCENT_GREEN, title, [desc], title_size=14, body_size=12)
    y += Inches(0.95)

# =============================================================================
# SLIDE 16 — SUMMARY
# =============================================================================
s = new_slide(); slide_header(s, "Summary", subtitle="Key Takeaways", tag="Closing"); footer(s)
card(s, Inches(0.45), Inches(1.2), Inches(12.4), Inches(3.6), ACCENT_GREEN, "What This Review Demonstrates",
     ["Every module proposed at the 0th review (M0-M5) is implemented, tested, and integrated into one",
      "live application -- not five disconnected scripts.",
      "",
      "The system has been run against real hardware and real live external APIs -- not simulated end to",
      "end -- and every result, positive or negative, is reported honestly.",
      "",
      "The security architecture (applicability -> verification -> authorization -> exploitation) is enforced",
      "in code and proven with live, adversarial-style negative tests, not just described.",
      "",
      "What remains is real-world scale: physical vulnerable hardware and edge deployment -- both are",
      "concretely scoped, not open questions."], body_size=14)

# =============================================================================
# SLIDE 17 — THANK YOU / Q&A
# =============================================================================
s = new_slide()
add_rect(s, 0, 0, SLIDE_W, Inches(0.1), ACCENT_GREEN)
add_rect(s, 0, 0, Inches(0.1), SLIDE_H, ACCENT_GREEN)
add_textbox(s, "Thank You", Inches(0.7), Inches(2.6), Inches(11.0), Inches(1.2), font_size=54, bold=True, color=WHITE)
add_textbox(s, "Questions & Discussion", Inches(0.7), Inches(3.75), Inches(11.0), Inches(0.6), font_size=22, color=ACCENT_GREEN)
add_textbox(s, "PISA -- Portable IoT Security Assessment Platform  |  Vivek Reddy  |  20CYS495",
            Inches(0.7), Inches(6.6), Inches(11.0), Inches(0.4), font_size=13, color=GRAY_LIGHT)

prs.save("docs/presentations/PISA_Review1.pptx")
print("Wrote docs/presentations/PISA_Review1.pptx —", len(prs.slides), "slides")
